#!/usr/bin/env python3
# SPDX-License-Identifier: GPL-3.0-or-later
"""Import validation data from a filled xlsx workbook into a model CSV.

Built for data/haiku_filled.xlsx (2026-04-06): the only surviving Haiku
validation data after the CSV-overwrite bug wiped the model CSVs. The
workbook uses the pre-2026-07 12-column layout (including the dropped
SINGLE_TOKEN_WORD? column) and encodes EXISTS? as '1'/'0'.

PROVENANCE WARNING (keep with any downstream use): the filled rows are
blanket per-root positives — all three vocalizations of a root marked
exists=1 with copy-pasted meanings (e.g. بَحُثَ marked real, which is
false) — and LIST_OF_DIALECTS holds free text. Import preserves the data;
it does not certify it.
"""

import csv
import sys
from pathlib import Path

from src.generate_csvs import (
    CSV_HEADERS,
    VALIDATION_HEADERS,
    assert_csv_schema,
    backup_csv,
)

# Columns we expect in the workbook header. SINGLE_TOKEN_WORD? is read but
# discarded (its values are uniformly garbage and the column is gone from
# the CSV schema).
REQUIRED_XLSX_COLUMNS = {'ROOT', 'PATTERN', 'VERB', 'EXISTS?'}

_BOOL_MAP = {
    '1': 'True', '0': 'False',
    'true': 'True', 'false': 'False',
    'yes': 'True', 'no': 'False',
}


def _norm_bool(value) -> str:
    """Map xlsx truthiness spellings to the CSV 'True'/'False' convention."""
    if value is None:
        return ''
    text = str(value).strip()
    if not text:
        return ''
    return _BOOL_MAP.get(text.lower(), text)


def read_filled_rows(xlsx_path: Path) -> list[dict]:
    """Rows from the workbook that have a non-empty EXISTS? cell,
    normalized to the current validation-column set."""
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else '' for h in next(rows_iter)]

    missing = REQUIRED_XLSX_COLUMNS - set(header)
    if missing:
        raise SystemExit(f"{xlsx_path} is missing expected columns: {sorted(missing)}")
    idx = {name: header.index(name) for name in header if name}

    filled = []
    for cells in rows_iter:
        if cells is None:
            continue

        def cell(name):
            i = idx.get(name)
            if i is None or i >= len(cells) or cells[i] is None:
                return ''
            return str(cells[i]).strip()

        if not cell('EXISTS?'):
            continue
        filled.append({
            'ROOT': cell('ROOT'),
            'PATTERN': cell('PATTERN'),
            'VERB': cell('VERB'),
            'EXISTS?': _norm_bool(cell('EXISTS?')),
            'DIALECT_ONLY?': _norm_bool(cell('DIALECT_ONLY?')),
            'LIST_OF_DIALECTS': cell('LIST_OF_DIALECTS'),
            'DIALECT_MEANINGS': cell('DIALECT_MEANINGS'),
            'MSA?': _norm_bool(cell('MSA?')),
            'MSA_MEANINGS': cell('MSA_MEANINGS'),
        })
    wb.close()
    return filled


def report(filled: list[dict], samples: int = 10):
    dist: dict[str, int] = {}
    for row in filled:
        dist[row['EXISTS?']] = dist.get(row['EXISTS?'], 0) + 1
    print(f"Filled rows found: {len(filled)}", file=sys.stderr)
    print(f"EXISTS? value distribution: {dist}", file=sys.stderr)
    print(f"Sample rows (first {samples}):", file=sys.stderr)
    for row in filled[:samples]:
        print(f"  {row['ROOT']} / {row['PATTERN']} / {row['VERB']} → "
              f"EXISTS?={row['EXISTS?']} MSA_MEANINGS={row['MSA_MEANINGS']!r}",
              file=sys.stderr)
    print(
        "\nPROVENANCE WARNING: rows are blanket per-root positives with "
        "copy-pasted meanings (quality suspect); LIST_OF_DIALECTS is free text.",
        file=sys.stderr,
    )


def import_into_csv(filled: list[dict], csv_path: Path) -> tuple[int, int]:
    """Apply filled rows to the target CSV by (ROOT, PATTERN).

    Returns (applied, unmatched). Backs the CSV up first; atomic write.
    """
    assert_csv_schema(csv_path)
    backup = backup_csv(csv_path, tag='preimport')
    print(f"Backup → {backup}", file=sys.stderr)

    by_key = {(row['ROOT'], row['PATTERN']): row for row in filled}
    matched_keys = set()

    with open(csv_path, 'r', encoding='utf-8', newline='') as f:
        rows = list(csv.DictReader(f))

    for row in rows:
        key = (row['ROOT'], row['PATTERN'])
        if key in by_key:
            src = by_key[key]
            for h in VALIDATION_HEADERS:
                row[h] = src.get(h, '')
            matched_keys.add(key)

    tmp_path = csv_path.with_suffix('.csv.tmp')
    with open(tmp_path, 'w', encoding='utf-8', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=CSV_HEADERS)
        writer.writeheader()
        writer.writerows(rows)
    tmp_path.rename(csv_path)

    unmatched = len(by_key) - len(matched_keys)
    if unmatched:
        print(f"WARNING: {unmatched} xlsx rows had no (ROOT, PATTERN) match "
              f"in {csv_path} and were not imported", file=sys.stderr)
    return len(matched_keys), unmatched
